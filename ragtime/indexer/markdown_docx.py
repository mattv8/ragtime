"""Bounded native Markdown to DOCX rendering for conversation exports."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from typing import Any
from urllib.parse import urlparse

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from markdown_it import MarkdownIt
from markdown_it.token import Token

from ragtime.core.logging import get_logger

logger = get_logger(__name__)

MAX_INPUT_BYTES = 2 * 1024 * 1024
MAX_TOKENS = 100_000
MAX_TABLE_CELLS = 10_000
MAX_NESTING = 32
_SAFE_LINK_SCHEMES = {"http", "https", "mailto"}
_TRIGGER_TYPES = {
    "heading_open",
    "bullet_list_open",
    "ordered_list_open",
    "table_open",
    "fence",
    "code_block",
    "blockquote_open",
    "hr",
    "strong_open",
    "em_open",
    "s_open",
    "code_inline",
    "link_open",
    "image",
}
_INVALID_XML_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")


class _BudgetExceeded(Exception):
    """Internal sentinel for a whole-document plain-text fallback."""


@dataclass
class _ListState:
    num_id: int
    level: int
    item_paragraph_seen: bool = False


def _clean(value: str) -> str:
    return _INVALID_XML_CHARS.sub("\ufffd", value)


def _safe_url(target: str) -> bool:
    if not target or any(character.isspace() or ord(character) < 32 for character in target):
        return False
    return urlparse(target).scheme.lower() in _SAFE_LINK_SCHEMES


def _append_numbering(document: Any, ordered: bool, level: int, start: int) -> int:
    numbering = document.part.numbering_part.element
    abstract_ids = [int(item.get(qn("w:abstractNumId"))) for item in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(item.get(qn("w:numId"))) for item in numbering.findall(qn("w:num"))]
    abstract_id, num_id = max(abstract_ids, default=0) + 1, max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi_level = OxmlElement("w:multiLevelType")
    multi_level.set(qn("w:val"), "multilevel")
    abstract.append(multi_level)
    for ilvl in range(9):
        level_element = OxmlElement("w:lvl")
        level_element.set(qn("w:ilvl"), str(ilvl))
        start_element = OxmlElement("w:start")
        start_element.set(qn("w:val"), "1")
        num_format = OxmlElement("w:numFmt")
        num_format.set(qn("w:val"), "decimal" if ordered else "bullet")
        level_text = OxmlElement("w:lvlText")
        level_text.set(qn("w:val"), f"%{ilvl + 1}." if ordered else "•")
        indent = OxmlElement("w:ind")
        indent.set(qn("w:left"), str(720 * (ilvl + 1)))
        indent.set(qn("w:hanging"), "360")
        paragraph_properties = OxmlElement("w:pPr")
        paragraph_properties.append(indent)
        level_element.extend((start_element, num_format, level_text, paragraph_properties))
        abstract.append(level_element)
    existing_nums = list(numbering.findall(qn("w:num")))
    numbering.insert(list(numbering).index(existing_nums[0]) if existing_nums else len(numbering), abstract)
    number = OxmlElement("w:num")
    number.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    number.append(abstract_ref)
    if ordered and start != 1:
        override = OxmlElement("w:lvlOverride")
        override.set(qn("w:ilvl"), str(level))
        start_override = OxmlElement("w:startOverride")
        start_override.set(qn("w:val"), str(start))
        override.append(start_override)
        number.append(override)
    numbering.append(number)
    return num_id


def _set_numbering(paragraph: Any, num_id: int, level: int) -> None:
    properties = paragraph._p.get_or_add_pPr()
    num_properties = OxmlElement("w:numPr")
    ilvl, number_id = OxmlElement("w:ilvl"), OxmlElement("w:numId")
    ilvl.set(qn("w:val"), str(level))
    number_id.set(qn("w:val"), str(num_id))
    num_properties.extend((ilvl, number_id))
    properties.append(num_properties)


def _add_run(paragraph: Any, text: str, state: tuple[int, int, int], code: bool = False, hyperlink: Any | None = None) -> None:
    run = paragraph.add_run(_clean(text))
    run.bold, run.italic, run.font.strike = bool(state[0]), bool(state[1]), bool(state[2])
    if code:
        run.font.name, run.font.size = "Courier New", Pt(9)
    if hyperlink is not None:
        hyperlink.append(run._r)


def _add_inline(paragraph: Any, tokens: list[Token], state: tuple[int, int, int] = (0, 0, 0), hyperlink: Any | None = None) -> None:
    current = state
    index = 0
    opening = {"strong_open": 0, "em_open": 1, "s_open": 2}
    closing = {"strong_close": 0, "em_close": 1, "s_close": 2}
    while index < len(tokens):
        token = tokens[index]
        if token.type in opening:
            position = opening[token.type]
            current = (current[0] + (position == 0), current[1] + (position == 1), current[2] + (position == 2))
        elif token.type in closing:
            position = closing[token.type]
            current = (max(0, current[0] - (position == 0)), max(0, current[1] - (position == 1)), max(0, current[2] - (position == 2)))
        elif token.type == "code_inline":
            _add_run(paragraph, token.content, current, code=True, hyperlink=hyperlink)
        elif token.type in {"softbreak", "hardbreak"}:
            run = paragraph.add_run()
            run.add_break()
            if hyperlink is not None:
                hyperlink.append(run._r)
        elif token.type == "link_open":
            close, depth = index + 1, 1
            while close < len(tokens) and depth:
                depth += tokens[close].type == "link_open"
                depth -= tokens[close].type == "link_close"
                close += 1
            children = tokens[index + 1 : close - 1]
            target = str(token.attrGet("href") or "")
            if _safe_url(target):
                relationship_id = paragraph.part.relate_to(
                    target, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True
                )
                link = OxmlElement("w:hyperlink")
                link.set(qn("r:id"), relationship_id)
                _add_inline(paragraph, children, current, link)
                paragraph._p.append(link)
            else:
                _add_inline(paragraph, children, current, hyperlink)
                if target:
                    _add_run(paragraph, f" ({target})", current, hyperlink=hyperlink)
            index = close - 1
        elif token.type == "image":
            alt, target = token.content or "", str(token.attrGet("src") or "")
            _add_run(paragraph, f"{alt} ({target})" if target else alt, current, hyperlink=hyperlink)
        elif token.type in {"text", "html_inline"} or token.content:
            _add_run(paragraph, token.content, current, hyperlink=hyperlink)
        index += 1


def _render(tokens: list[Token], document: Any) -> None:
    lists: list[_ListState] = []
    quote_depth = 0
    index = 0
    while index < len(tokens):
        token = tokens[index]
        if token.type in {"bullet_list_open", "ordered_list_open"}:
            if len(lists) >= 9:
                raise _BudgetExceeded
            start = int(token.attrGet("start") or "1")
            lists.append(_ListState(_append_numbering(document, token.type == "ordered_list_open", len(lists), start), len(lists)))
        elif token.type in {"bullet_list_close", "ordered_list_close"}:
            lists.pop()
        elif token.type == "list_item_open":
            lists[-1].item_paragraph_seen = False
        elif token.type == "blockquote_open":
            quote_depth += 1
        elif token.type == "blockquote_close":
            quote_depth -= 1
        elif token.type == "heading_open" and index + 1 < len(tokens):
            paragraph = document.add_heading(level=int(token.tag[1]))
            _add_inline(paragraph, tokens[index + 1].children or [])
            index += 2
        elif token.type == "paragraph_open" and index + 1 < len(tokens):
            paragraph = document.add_paragraph()
            if lists:
                current = lists[-1]
                if not current.item_paragraph_seen:
                    _set_numbering(paragraph, current.num_id, current.level)
                    current.item_paragraph_seen = True
                else:
                    paragraph.paragraph_format.left_indent = Inches(0.5 * (current.level + 1))
            elif quote_depth:
                paragraph.paragraph_format.left_indent = Inches(0.35 * quote_depth)
            _add_inline(paragraph, tokens[index + 1].children or [])
            index += 2
        elif token.type in {"fence", "code_block"}:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.25)
            _add_run(paragraph, token.content, (0, 0, 0), code=True)
        elif token.type == "hr":
            document.add_paragraph("─" * 48)
        elif token.type == "table_open":
            end = next(position for position in range(index, len(tokens)) if tokens[position].type == "table_close")
            rows: list[list[list[Token]]] = []
            header_rows: set[int] = set()
            row: list[list[Token]] | None = None
            for table_token in tokens[index + 1 : end]:
                if table_token.type == "tr_open":
                    row = []
                    rows.append(row)
                elif table_token.type in {"th_open", "td_open"} and row is not None:
                    row.append([])
                    if table_token.type == "th_open":
                        header_rows.add(len(rows) - 1)
                elif table_token.type == "inline" and row:
                    row[-1] = table_token.children or []
            if rows:
                table = document.add_table(rows=len(rows), cols=max(len(row) for row in rows))
                table.style = "Table Grid"
                for row_index, cells in enumerate(rows):
                    for cell_index, children in enumerate(cells):
                        paragraph = table.cell(row_index, cell_index).paragraphs[0]
                        _add_inline(paragraph, children, (1 if row_index in header_rows else 0, 0, 0))
            index = end
        elif token.content:
            document.add_paragraph(_clean(token.content))
        index += 1


def render_markdown_docx(text: str, title: str) -> bytes | None:
    """Return native styled DOCX bytes, or ``None`` for a safe plain-text fallback."""
    try:
        input_bytes = len(text.encode("utf-8", "replace"))
        if input_bytes > MAX_INPUT_BYTES:
            logger.info("Markdown DOCX conversion skipped: input_bytes=%d", input_bytes)
            return None
        parser = (
            MarkdownIt("commonmark", {"maxNesting": MAX_NESTING, "html": False, "linkify": False, "typographer": False}).enable("table").enable("strikethrough")
        )
        tokens = parser.parse(_clean(text))
        all_tokens = [item for block_token in tokens for item in (block_token, *(block_token.children or []))]
        token_count = len(all_tokens)
        table_cells = sum(token.type in {"th_open", "td_open"} for token in all_tokens)
        if token_count > MAX_TOKENS or table_cells > MAX_TABLE_CELLS or any(token.level >= MAX_NESTING - 2 for token in all_tokens):
            logger.info("Markdown DOCX conversion skipped: tokens=%d table_cells=%d", token_count, table_cells)
            return None
        if not any(token.type in _TRIGGER_TYPES for token in all_tokens):
            return None
        document = Document()
        if title:
            document.add_heading(_clean(title), level=1)
        _render(tokens, document)
        output = io.BytesIO()
        document.save(output)
        return output.getvalue()
    except _BudgetExceeded:
        logger.info("Markdown DOCX conversion skipped: nesting_limit")
        return None
    except Exception as error:
        logger.warning("Markdown DOCX conversion failed: exception=%s", type(error).__name__)
        return None
