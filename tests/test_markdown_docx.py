from __future__ import annotations

from io import BytesIO
from typing import Any
from unittest import TestCase, mock
from zipfile import ZipFile

from docx import Document

from ragtime.indexer.markdown_docx import render_markdown_docx


class MarkdownDocxTests(TestCase):
    def document(self, markdown: str) -> Any:
        rendered = render_markdown_docx(markdown, "Export title")
        self.assertIsNotNone(rendered)
        return Document(BytesIO(rendered))  # type: ignore[arg-type]

    def test_formats_headings_and_nested_inline_runs(self) -> None:
        document = self.document("# Heading\n\n**bold and *italic***, ~~gone~~, and `a_b()`")

        self.assertEqual(document.paragraphs[1].style.name, "Heading 1")
        runs = document.paragraphs[2].runs
        self.assertTrue(any(run.bold and run.text == "bold and " for run in runs))
        self.assertTrue(any(run.bold and run.italic and run.text == "italic" for run in runs))
        self.assertTrue(any(run.font.strike and run.text == "gone" for run in runs))
        self.assertTrue(any(run.font.name == "Courier New" and run.text == "a_b()" for run in runs))

    def test_preserves_escaped_code_newlines_html_comments_and_ascii_tables(self) -> None:
        document = self.document("# heading\n\n\\*literal\\*\n\n```python\nif a < b:\n\tprint('x')\n```\n\n<!-- note -->\na----+----b")

        content = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("*literal*", content)
        self.assertIn("if a < b:\n\tprint('x')\n", content)
        self.assertIn("<!-- note -->", content)
        self.assertIn("a----+----b", content)
        self.assertEqual(document.paragraphs[3].runs[0].font.name, "Courier New")

    def test_native_numbering_order_levels_restarts_and_continuations(self) -> None:
        markdown = "3. three\n\n   continuation\n\n   * child\n      1. grandchild\n4. four\n\n---\n\n1. restart"
        rendered = render_markdown_docx(markdown, "Export title")
        self.assertIsNotNone(rendered)
        document = Document(BytesIO(rendered))  # type: ignore[arg-type]
        numbered = [paragraph for paragraph in document.paragraphs if "numPr" in paragraph._p.xml]
        self.assertEqual(len(numbered), 5)
        self.assertNotIn("numPr", document.paragraphs[2]._p.xml)
        with ZipFile(BytesIO(rendered)) as archive:  # type: ignore[arg-type]
            numbering = archive.read("word/numbering.xml").decode("utf-8")
        custom_abstract = numbering.rindex("<w:abstractNum w:abstractNumId=")
        custom_num = numbering.rindex("<w:num w:numId=")
        self.assertLess(custom_abstract, custom_num)
        self.assertIn('<w:multiLevelType w:val="multilevel"', numbering)
        self.assertIn('<w:startOverride w:val="3"', numbering)
        self.assertIn('<w:ilvl w:val="2"', document.part.element.xml)

    def test_nested_ordered_list_start_override_and_separate_restart_use_fresh_ids(self) -> None:
        rendered = render_markdown_docx("1. outer\n\n   4. nested\n\n---\n\n1. restarted", "Export title")
        self.assertIsNotNone(rendered)
        document = Document(BytesIO(rendered))  # type: ignore[arg-type]
        numbered = [paragraph for paragraph in document.paragraphs if "numPr" in paragraph._p.xml]
        num_ids = [paragraph._p.xpath("./w:pPr/w:numPr/w:numId/@w:val")[0] for paragraph in numbered]
        self.assertEqual(len(num_ids), 3)
        self.assertNotEqual(num_ids[0], num_ids[1])
        self.assertNotEqual(num_ids[0], num_ids[2])
        with ZipFile(BytesIO(rendered)) as archive:  # type: ignore[arg-type]
            numbering = archive.read("word/numbering.xml").decode("utf-8")
        self.assertIn('<w:lvlOverride w:ilvl="1"><w:startOverride w:val="4"', numbering)

    def test_safe_links_preserve_nested_formatting_and_unsafe_targets_stay_visible(self) -> None:
        rendered = render_markdown_docx("[**bold `code`**](https://example.test) [bad](javascript:alert(1)) [anchor](#part)", "title")
        self.assertIsNotNone(rendered)
        with ZipFile(BytesIO(rendered)) as archive:  # type: ignore[arg-type]
            document_xml = archive.read("word/document.xml")
            relationships = archive.read("word/_rels/document.xml.rels")
        self.assertIn(b"https://example.test", relationships)
        self.assertIn(b"<w:b/>", document_xml)
        self.assertIn(b"Courier New", document_xml)
        self.assertIn(b"javascript:alert(1)", document_xml)
        self.assertIn(b"(#part)", document_xml)

    def test_images_are_literal_and_never_create_media_parts(self) -> None:
        rendered = render_markdown_docx("# title\n\n![chart](https://image.test/chart.png)", "title")
        self.assertIsNotNone(rendered)
        with ZipFile(BytesIO(rendered)) as archive:  # type: ignore[arg-type]
            self.assertNotIn("word/media/", "\n".join(archive.namelist()))
        self.assertIn("chart (https://image.test/chart.png)", self.document("# title\n\n![chart](https://image.test/chart.png)").paragraphs[-1].text)

    def test_tables_quotes_and_xml_sanitization(self) -> None:
        document = self.document("> quoted\n\n| **Header** | Value |\n| --- | --- |\n| `cell` | value |\n\n# valid\x01\ud800\ufffe")

        self.assertEqual(len(document.tables), 1)
        self.assertTrue(document.tables[0].cell(0, 0).paragraphs[0].runs[0].bold)
        self.assertEqual(document.tables[0].cell(1, 0).paragraphs[0].runs[0].font.name, "Courier New")
        self.assertIn("quoted", document.paragraphs[1].text)
        self.assertIn("valid\ufffd\ufffd\ufffd", document.paragraphs[-1].text)

    def test_plain_deep_and_actual_parser_budgets_fall_back_with_sentinel(self) -> None:
        self.assertIsNone(render_markdown_docx("ordinary_text with punctuation", "title"))
        sentinel = "LAST-SENTINEL"
        self.assertIsNone(render_markdown_docx("> " * 32 + sentinel, "title"))
        self.assertIsNone(render_markdown_docx("# " + "x" * (2 * 1024 * 1024), "title"))
        self.assertIsNone(render_markdown_docx("# heading\n\n" + "*x* " * 50_001, "title"))
        table = "| h |\n| --- |\n" + "| x |\n" * 10_001
        self.assertIsNone(render_markdown_docx(table, "title"))

    def test_deeper_than_word_levels_and_document_failure_fall_back(self) -> None:
        nested = "\n".join("  " * level + "* item" for level in range(10))
        self.assertIsNone(render_markdown_docx(nested, "title"))
        with mock.patch("ragtime.indexer.markdown_docx.Document", side_effect=RuntimeError):
            self.assertIsNone(render_markdown_docx("# heading", "title"))

    def test_budget_and_failure_logs_exclude_document_and_exception_text(self) -> None:
        source_sentinel = "SOURCE-UNIQUE-SENTINEL"
        title_sentinel = "TITLE-UNIQUE-SENTINEL"
        with self.assertLogs("ragtime.indexer.markdown_docx", level="INFO") as budget_logs:
            self.assertIsNone(render_markdown_docx("# " + source_sentinel * 110_000, title_sentinel))
        with mock.patch("ragtime.indexer.markdown_docx.Document", side_effect=RuntimeError("EXCEPTION-UNIQUE-SENTINEL")):
            with self.assertLogs("ragtime.indexer.markdown_docx", level="WARNING") as failure_logs:
                self.assertIsNone(render_markdown_docx("# " + source_sentinel, title_sentinel))
        logged = "\n".join(budget_logs.output + failure_logs.output)
        self.assertNotIn(source_sentinel, logged)
        self.assertNotIn(title_sentinel, logged)
        self.assertNotIn("EXCEPTION-UNIQUE-SENTINEL", logged)
