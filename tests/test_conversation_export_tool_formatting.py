from __future__ import annotations

import asyncio
import base64
import io
import json
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Any, Coroutine, cast

from docx import Document

from ragtime.indexer import export_service
from ragtime.rag.components import RAGComponents


class ConversationExportToolFormattingTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmpdir = TemporaryDirectory()
        self.original_base_dir = export_service.EXPORT_BASE_DIR
        export_service.EXPORT_BASE_DIR = Path(self.tmpdir.name) / "exports"

    def tearDown(self) -> None:
        export_service.EXPORT_BASE_DIR = self.original_base_dir
        self.tmpdir.cleanup()

    def _create_with_tool(self, **kwargs: Any) -> tuple[dict[str, Any], dict[str, Any]]:
        tool = RAGComponents()._build_conversation_export_tool(
            conversation_id="conversation-1",
            workspace_id="workspace-1",
        )
        assert tool is not None and tool.coroutine is not None

        export_dir = export_service.EXPORT_BASE_DIR / "conversation-1"
        existing_paths = set(export_dir.glob("*.json")) if export_dir.exists() else set()
        coroutine = cast(Coroutine[Any, Any, str], tool.coroutine(**kwargs))
        payload = json.loads(asyncio.run(coroutine))
        new_paths = set(export_dir.glob("*.json")) - existing_paths
        self.assertEqual(len(new_paths), 1)
        spec_path = new_paths.pop()
        return payload, json.loads(spec_path.read_text(encoding="utf-8"))

    def test_markdown_docx_is_prepared_before_download_and_reported_without_leaking_content(self) -> None:
        markdown = "# Quarterly report\n\n- First result\n- **Second result**"

        payload, spec = self._create_with_tool(
            filename="quarterly-report.docx",
            format="docx",
            text=markdown,
        )

        prepared = spec["prepared_docx"]
        document = Document(io.BytesIO(base64.b64decode(prepared["content_base64"], validate=True)))
        markdown_heading = next(paragraph for paragraph in document.paragraphs if paragraph.text == "Quarterly report")
        heading_style = markdown_heading.style
        assert heading_style is not None
        self.assertEqual(heading_style.name, "Heading 1")
        self.assertNotIn("#", markdown_heading.text)
        list_paragraphs = [paragraph for paragraph in document.paragraphs if paragraph.text in {"First result", "Second result"}]
        self.assertEqual(len(list_paragraphs), 2)
        for paragraph in list_paragraphs:
            self.assertIsNotNone(paragraph._p.pPr)
            assert paragraph._p.pPr is not None
            self.assertIsNotNone(paragraph._p.pPr.numPr)
        second_run = next(run for run in list_paragraphs[1].runs if run.text == "Second result")
        self.assertTrue(second_run.bold)
        self.assertTrue(payload["markdown_formatted"])
        self.assertEqual(spec["source"]["kind"], "content_snapshot")
        self.assertEqual(spec["source"]["text"], markdown)
        self.assertEqual(
            set(payload),
            {
                "tool",
                "status",
                "filename",
                "format",
                "source_kind",
                "reused_previous_source",
                "download_url",
                "markdown_link",
                "expires_at",
                "instruction",
                "markdown_formatted",
            },
        )
        self.assertNotIn("prepared_docx", payload)
        self.assertNotIn(markdown, json.dumps(payload))
        self.assertNotIn(prepared["content_base64"], json.dumps(payload))

    def test_oversized_markdown_docx_is_not_reported_as_formatted_and_preserves_source(self) -> None:
        markdown = "# " + ("x" * (2 * 1024 * 1024))

        payload, spec = self._create_with_tool(
            filename="oversized.docx",
            format="docx",
            text=markdown,
        )

        self.assertFalse(payload["markdown_formatted"])
        self.assertNotIn("prepared_docx", spec)
        self.assertEqual(spec["source"]["kind"], "content_snapshot")
        self.assertEqual(spec["source"]["text"], markdown)

    def test_plain_markdown_download_and_binary_docx_are_not_reported_as_markdown_formatted(self) -> None:
        binary_document = Document()
        binary_document.add_paragraph("Prebuilt document")
        binary_buffer = io.BytesIO()
        binary_document.save(binary_buffer)
        binary_docx = binary_buffer.getvalue()

        cases = (
            ("plain.docx", "docx", "Plain document text.", "", False),
            ("notes.md", "md", "# Intentional Markdown download", "", False),
            ("prebuilt.docx", "docx", "", base64.b64encode(binary_docx).decode("ascii"), True),
        )
        for filename, export_format, text, content_base64, is_binary in cases:
            with self.subTest(filename=filename):
                payload, spec = self._create_with_tool(
                    filename=filename,
                    format=export_format,
                    text=text,
                    content_base64=content_base64,
                    mime_type=("application/vnd.openxmlformats-officedocument.wordprocessingml.document" if is_binary else ""),
                )

                self.assertFalse(payload["markdown_formatted"])
                self.assertNotIn("prepared_docx", spec)
                if is_binary:
                    self.assertEqual(spec["source"]["kind"], "binary_snapshot")
                    self.assertEqual(base64.b64decode(spec["source"]["content_base64"], validate=True), binary_docx)
                else:
                    self.assertEqual(spec["source"]["kind"], "content_snapshot")
                    self.assertEqual(spec["source"]["text"], text)
