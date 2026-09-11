from __future__ import annotations

import asyncio
import base64
import hashlib
import hmac
import io
import json
import unittest
from datetime import datetime, timezone
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Any, Coroutine, cast
from unittest import mock

from docx import Document
from docx.text.paragraph import Paragraph
from langchain_core.messages import AIMessage, ToolMessage
from openpyxl import load_workbook

from ragtime.core.sql_utils import format_query_result
from ragtime.indexer import export_service
from ragtime.indexer.models import CreateConversationExportRequest
from ragtime.rag.components import RAGComponents
from ragtime.tools.datatable import create_datatable


def _paragraph_style_name(paragraph: Paragraph) -> str:
    style = paragraph.style
    assert style is not None
    name = style.name
    assert name is not None
    return name


class ConversationExportServiceTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmpdir = TemporaryDirectory()
        self.original_base_dir = export_service.EXPORT_BASE_DIR
        export_service.EXPORT_BASE_DIR = Path(self.tmpdir.name) / "exports"

    def tearDown(self) -> None:
        export_service.EXPORT_BASE_DIR = self.original_base_dir
        self.tmpdir.cleanup()

    def test_create_export_spec_sanitizes_filename_and_builds_markdown_url(self) -> None:
        spec = export_service.create_export_spec(
            conversation_id="conversation-1",
            filename="../Revenue Report.xlsx",
            export_format="xlsx",
            source=export_service.table_source(["Name"], [["Acme"]]),
            workspace_id="workspace-1",
            title="Revenue Report",
        )

        self.assertEqual(spec["filename"], "Revenue_Report.xlsx")
        self.assertIn("/indexes/conversations/conversation-1/exports/", spec["download_url"])
        self.assertIn("Revenue_Report.xlsx", spec["download_url"])
        self.assertIn("workspace_id=workspace-1", spec["download_url"])
        export_service.verify_token(spec["token"], "conversation-1", spec["id"], spec["filename"])

    def test_export_links_do_not_expire(self) -> None:
        expired_at = datetime(2000, 1, 1, tzinfo=timezone.utc)
        spec = export_service.create_export_spec(
            conversation_id="conversation-1",
            filename="old.csv",
            export_format="csv",
            source=export_service.table_source(["A"], [[1]]),
            expires_in_seconds=60,
        )
        path = export_service.EXPORT_BASE_DIR / "conversation-1" / f"{spec['id']}.json"
        payload = json.loads(path.read_text(encoding="utf-8"))
        payload["expires_at"] = expired_at.isoformat()
        path.write_text(json.dumps(payload), encoding="utf-8")
        token = export_service.create_token("conversation-1", spec["id"], spec["filename"], expired_at)

        export_service.verify_token(token, "conversation-1", spec["id"], spec["filename"])
        loaded_spec = export_service.load_export_spec("conversation-1", spec["id"])

        self.assertEqual(loaded_spec["id"], spec["id"])
        self.assertTrue(path.exists())

    def test_verify_token_accepts_legacy_exp_payload(self) -> None:
        expired_at = datetime(2000, 1, 1, tzinfo=timezone.utc)
        payload = {
            "conversation_id": "conversation-1",
            "export_id": "export-1",
            "filename": "old.csv",
            "exp": int(expired_at.timestamp()),
        }
        payload_b64 = export_service._base64_url_encode(json.dumps(payload, sort_keys=True, separators=(",", ":")).encode("utf-8"))
        signature = hmac.new(
            export_service._secret(),
            payload_b64.encode("ascii"),
            hashlib.sha256,
        ).digest()
        token = f"{payload_b64}.{export_service._base64_url_encode(signature)}"

        verified = export_service.verify_token(token, "conversation-1", "export-1", "old.csv")

        self.assertEqual(verified["exp"], int(expired_at.timestamp()))

    def test_render_table_formats_escape_formula_cells(self) -> None:
        spec = export_service.create_export_spec(
            conversation_id="conversation-1",
            filename="accounts.csv",
            export_format="csv",
            source=export_service.table_source(["Name", "Formula"], [["Acme", "=2+2"]]),
        )

        data, media_type = asyncio.run(export_service.render_export(spec))

        self.assertEqual(media_type, export_service.MIME_TYPES["csv"])
        self.assertIn("Acme", data.decode("utf-8-sig"))
        self.assertIn("\t=2+2", data.decode("utf-8-sig"))

    def test_render_xlsx_table(self) -> None:
        spec = export_service.create_export_spec(
            conversation_id="conversation-1",
            filename="accounts.xlsx",
            export_format="xlsx",
            source=export_service.table_source(["Name", "Total"], [["Acme", 42]]),
        )

        data, media_type = asyncio.run(export_service.render_export(spec))
        path = Path(self.tmpdir.name) / "accounts.xlsx"
        path.write_bytes(data)
        workbook = load_workbook(path)

        self.assertEqual(media_type, export_service.MIME_TYPES["xlsx"])
        sheet = workbook.active
        assert sheet is not None
        self.assertEqual(sheet["A1"].value, "Name")
        self.assertEqual(sheet["A2"].value, "Acme")
        self.assertEqual(sheet["B2"].value, 42)

    def test_render_document_formats_from_content(self) -> None:
        for export_format, signature in (("pdf", b"%PDF-"), ("docx", b"PK")):
            with self.subTest(export_format=export_format):
                spec = export_service.create_export_spec(
                    conversation_id="conversation-1",
                    filename=f"summary.{export_format}",
                    export_format=export_format,
                    source=export_service.content_source(text="Quarterly summary"),
                    title="Summary",
                )

                data, media_type = asyncio.run(export_service.render_export(spec))

                self.assertTrue(data.startswith(signature))
                self.assertEqual(media_type, export_service.MIME_TYPES[export_format])

    def test_saved_markdown_docx_uses_native_heading_formatting(self) -> None:
        spec = export_service.create_export_spec(
            conversation_id="conversation-1",
            filename="formatted.docx",
            export_format="docx",
            source=export_service.content_source(text="# Quarterly summary\n\n- Revenue increased"),
            title="Summary",
        )
        spec_path = export_service.EXPORT_BASE_DIR / "conversation-1" / f"{spec['id']}.json"
        saved_spec_bytes = spec_path.read_bytes()
        loaded = export_service.load_export_spec("conversation-1", spec["id"])

        prepared = cast(dict[str, Any], loaded["prepared_docx"])
        prepared_bytes = base64.b64decode(prepared["content_base64"], validate=True)
        prepared_document = Document(io.BytesIO(prepared_bytes))

        data, media_type = asyncio.run(export_service.render_export(loaded))
        document = Document(io.BytesIO(data))
        reloaded = export_service.load_export_spec("conversation-1", spec["id"])

        self.assertEqual(media_type, export_service.MIME_TYPES["docx"])
        self.assertEqual(data, prepared_bytes)
        self.assertEqual(spec_path.read_bytes(), saved_spec_bytes)
        self.assertEqual(reloaded["source"], spec["source"])
        self.assertIn("Heading", _paragraph_style_name(prepared_document.paragraphs[1]))
        self.assertIn("Heading", _paragraph_style_name(document.paragraphs[1]))

    def test_prepared_docx_is_persisted_and_reused_without_download_mutation(self) -> None:
        source = export_service.content_source(text="# Prepared heading\n\n- Prepared item")
        spec = export_service.create_export_spec(
            conversation_id="conversation-1", filename="prepared.docx", export_format="docx", source=source, title="Prepared"
        )
        path = export_service.EXPORT_BASE_DIR / "conversation-1" / f"{spec['id']}.json"
        saved_bytes = path.read_bytes()
        prepared = cast(dict[str, Any], spec["prepared_docx"])
        prepared_bytes = base64.b64decode(prepared["content_base64"], validate=True)

        first, _ = asyncio.run(export_service.render_export(export_service.load_export_spec("conversation-1", spec["id"])))
        second, _ = asyncio.run(export_service.render_export(export_service.load_export_spec("conversation-1", spec["id"])))

        self.assertEqual(source, spec["source"])
        self.assertEqual(first, prepared_bytes)
        self.assertEqual(second, prepared_bytes)
        self.assertEqual(path.read_bytes(), saved_bytes)

    def test_legacy_stale_and_malformed_prepared_docx_fall_back_to_renderer(self) -> None:
        spec = export_service.create_export_spec(
            conversation_id="conversation-1",
            filename="fallback.docx",
            export_format="docx",
            source=export_service.content_source(text="# Fresh heading"),
            title="Fresh",
        )
        for mutation in (
            lambda current: current.pop("prepared_docx"),
            lambda current: current.update(title="Changed title"),
            lambda current: current["source"].update(text="# Changed heading"),
            lambda current: current["prepared_docx"].update(content_base64="not base64"),
            lambda current: current["prepared_docx"].update(version=2),
            lambda current: current["prepared_docx"].update(source_sha256="not-ascii-\u00e9"),
            lambda current: current["prepared_docx"].update(content_base64=base64.b64encode(b"not a docx").decode("ascii")),
            lambda current: current["prepared_docx"].update(content_base64="A" * (export_service._PREPARED_DOCX_MAX_BASE64_CHARS + 1)),
        ):
            with self.subTest(mutation=mutation):
                current = json.loads(json.dumps(spec))
                mutation(current)

                data, media_type = asyncio.run(export_service.render_export(current))

                self.assertEqual(media_type, export_service.MIME_TYPES["docx"])
                self.assertTrue(data.startswith(b"PK"))
                self.assertIn("Heading", _paragraph_style_name(Document(io.BytesIO(data)).paragraphs[1]))

    def test_only_markdown_content_docx_gets_prepared_metadata(self) -> None:
        cases = (
            ("txt", export_service.content_source(text="# Markdown")),
            ("docx", export_service.content_source(text="plain text")),
            ("docx", export_service.content_source(content_base64=base64.b64encode(b"binary").decode("ascii"))),
            ("docx", export_service.table_source(["Column"], [["# Markdown"]])),
        )
        for export_format, source in cases:
            with self.subTest(export_format=export_format, source_kind=source["kind"]):
                spec = export_service.create_export_spec(
                    conversation_id="conversation-1", filename=f"no-prep.{export_format}", export_format=export_format, source=source
                )
                self.assertNotIn("prepared_docx", spec)

    def test_create_request_ignores_client_prepared_docx_metadata(self) -> None:
        request = CreateConversationExportRequest.model_validate(
            {
                "filename": "client.docx",
                "format": "docx",
                "source_kind": "content",
                "text": "# Server generated",
                "prepared_docx": {"version": 1, "content_base64": "client bytes"},
            }
        )

        self.assertNotIn("prepared_docx", request.model_dump())

    def test_preparation_failure_preserves_the_original_source(self) -> None:
        source = export_service.content_source(text="# Safe fallback")
        with mock.patch.object(export_service, "render_markdown_docx", side_effect=RuntimeError("sensitive source text")):
            spec = export_service.create_export_spec(
                conversation_id="conversation-1", filename="failure.docx", export_format="docx", source=source, title="Failure"
            )

        self.assertEqual(spec["source"], source)
        self.assertNotIn("prepared_docx", spec)

    def test_preparation_packaging_failure_preserves_the_original_source(self) -> None:
        source = export_service.content_source(text="# Safe packaging fallback")
        original_b64encode = export_service.base64.b64encode
        calls = 0

        def fail_preparation_only(value: bytes) -> bytes:
            nonlocal calls
            calls += 1
            if calls == 1:
                raise RuntimeError("sensitive source text")
            return original_b64encode(value)

        with mock.patch.object(export_service.base64, "b64encode", side_effect=fail_preparation_only):
            spec = export_service.create_export_spec(
                conversation_id="conversation-1", filename="packaging.docx", export_format="docx", source=source, title="Packaging"
            )

        saved = export_service.load_export_spec("conversation-1", spec["id"])
        self.assertEqual(spec["source"], source)
        self.assertEqual(saved["source"], source)
        self.assertNotIn("prepared_docx", spec)

    def test_plain_content_docx_preserves_paragraphs_semantically(self) -> None:
        spec = export_service.create_export_spec(
            conversation_id="conversation-1",
            filename="plain.docx",
            export_format="docx",
            source=export_service.content_source(text="First paragraph\nSecond paragraph"),
            title="Plain export",
        )

        data, media_type = asyncio.run(export_service.render_export(spec))
        document = Document(io.BytesIO(data))

        self.assertEqual(media_type, export_service.MIME_TYPES["docx"])
        self.assertEqual([paragraph.text for paragraph in document.paragraphs], ["Plain export", "First paragraph", "Second paragraph"])

    def test_plain_punctuation_and_escaped_markdown_preserve_original_paragraphs(self) -> None:
        text = "Budget_2026 and 3*5\n\\# Escaped heading\n\\- Escaped list"
        spec = export_service.create_export_spec(
            conversation_id="conversation-1",
            filename="literal.docx",
            export_format="docx",
            source=export_service.content_source(text=text),
            title="Literal export",
        )

        data, media_type = asyncio.run(export_service.render_export(spec))
        document = Document(io.BytesIO(data))

        self.assertEqual(media_type, export_service.MIME_TYPES["docx"])
        self.assertEqual([paragraph.text for paragraph in document.paragraphs], ["Literal export", *text.split("\n")])

    def test_content_markdown_and_legacy_doc_bytes_are_unchanged(self) -> None:
        text = "# Literal markdown\n\n- Item"
        for export_format in ("md", "doc"):
            with self.subTest(export_format=export_format):
                spec = export_service.create_export_spec(
                    conversation_id="conversation-1",
                    filename=f"literal.{export_format}",
                    export_format=export_format,
                    source=export_service.content_source(text=text),
                    title="Literal",
                )

                data, media_type = asyncio.run(export_service.render_export(spec))

                self.assertEqual(media_type, export_service.MIME_TYPES[export_format])
                if export_format == "md":
                    self.assertEqual(data, text.encode("utf-8"))
                else:
                    self.assertEqual(
                        data,
                        b'<!doctype html><html><head><meta charset="utf-8"></head><body><pre># Literal markdown\n\n- Item</pre></body></html>',
                    )

    def test_binary_docx_passthrough_is_byte_identical(self) -> None:
        source_document = Document()
        source_paragraph = source_document.add_paragraph("# Literal Markdown ")
        source_paragraph.add_run("bold source text").bold = True
        source_buffer = io.BytesIO()
        source_document.save(source_buffer)
        source_bytes = source_buffer.getvalue()
        spec = export_service.create_export_spec(
            conversation_id="conversation-1",
            filename="uploaded.docx",
            export_format="docx",
            source=export_service.content_source(
                content_base64=base64.b64encode(source_bytes).decode("ascii"),
                mime_type=export_service.MIME_TYPES["docx"],
            ),
        )

        data, media_type = asyncio.run(export_service.render_export(spec))

        self.assertEqual(data, source_bytes)
        self.assertEqual(media_type, export_service.MIME_TYPES["docx"])

        document = Document(io.BytesIO(data))
        self.assertEqual(document.paragraphs[0].text, "# Literal Markdown bold source text")
        self.assertTrue(document.paragraphs[0].runs[1].bold)

    def test_table_docx_contents_are_unchanged(self) -> None:
        spec = export_service.create_export_spec(
            conversation_id="conversation-1",
            filename="accounts.docx",
            export_format="docx",
            source=export_service.table_source(["Name", "Total"], [["Acme", 42], ["Globex", 7]]),
            title="Accounts",
        )

        data, media_type = asyncio.run(export_service.render_export(spec))
        document = Document(io.BytesIO(data))

        self.assertEqual(media_type, export_service.MIME_TYPES["docx"])
        self.assertEqual(
            [[cell.text for cell in row.cells] for row in document.tables[0].rows],
            [["Name", "Total"], ["Acme", "42"], ["Globex", "7"]],
        )

    def test_sizable_markdown_document_renders_headings_and_lists(self) -> None:
        text = "\n\n".join(f"## Section {index}\n\n" + "\n".join(f"- Item {index}-{item}" for item in range(6)) for index in range(78))
        spec = export_service.create_export_spec(
            conversation_id="conversation-1",
            filename="large.docx",
            export_format="docx",
            source=export_service.content_source(text=text),
            title="Large report",
        )

        data, _ = asyncio.run(export_service.render_export(spec))
        document = Document(io.BytesIO(data))

        self.assertEqual(sum(_paragraph_style_name(paragraph).startswith("Heading") for paragraph in document.paragraphs), 79)
        list_paragraphs = [paragraph for paragraph in document.paragraphs if paragraph.text.startswith("Item ")]
        self.assertEqual(len(list_paragraphs), 78 * 6)
        self.assertTrue(all(paragraph._p.pPr is not None and paragraph._p.pPr.numPr is not None for paragraph in list_paragraphs))

    def test_deep_markdown_falls_back_without_losing_final_sentinel(self) -> None:
        text = f"{'> ' * 32}Deep content\n\nFINAL_SENTINEL"
        spec = export_service.create_export_spec(
            conversation_id="conversation-1",
            filename="deep.docx",
            export_format="docx",
            source=export_service.content_source(text=text),
            title="Deep export",
        )

        data, media_type = asyncio.run(export_service.render_export(spec))
        document = Document(io.BytesIO(data))

        self.assertEqual(media_type, export_service.MIME_TYPES["docx"])
        self.assertEqual([paragraph.text for paragraph in document.paragraphs], ["Deep export", *text.split("\n")])

    def test_render_binary_snapshot_for_common_filetype(self) -> None:
        content = b"custom-binary-content"
        spec = export_service.create_export_spec(
            conversation_id="conversation-1",
            filename="archive.zip",
            export_format="zip",
            source=export_service.content_source(
                content_base64=base64.b64encode(content).decode("ascii"),
                mime_type="application/zip",
            ),
            mime_type="application/zip",
        )

        data, media_type = asyncio.run(export_service.render_export(spec))

        self.assertEqual(data, content)
        self.assertEqual(media_type, "application/zip")

    def test_live_table_export_uses_click_time_resolver(self) -> None:
        spec = export_service.create_export_spec(
            conversation_id="conversation-1",
            filename="live.csv",
            export_format="csv",
            source=export_service.live_table_source(
                {"component_id": "tool-1", "request": {"query": "select * from accounts"}},
                columns=["stale"],
                rows=[["old"]],
            ),
        )
        captured_sources: list[dict[str, object]] = []

        async def resolver(source: dict[str, object]):
            captured_sources.append(source)
            return ["Name"], [["Fresh"]]

        data, _ = asyncio.run(export_service.render_export(spec, live_table_resolver=resolver))
        payload = data.decode("utf-8-sig")

        self.assertEqual(len(captured_sources), 1)
        self.assertIn("Fresh", payload)
        self.assertNotIn("old", payload)

    def test_extracts_latest_export_context_from_datatable_output(self) -> None:
        data_connection = {
            "component_kind": "tool_config",
            "component_id": "tool-1",
            "request": {"query": "select name, total from accounts"},
        }
        output = asyncio.run(
            create_datatable(
                title="Accounts",
                columns=["Name", "Total"],
                data=[["Acme", 42]],
                data_connection=data_connection,
            )
        )

        context = RAGComponents()._build_export_context_from_visualization_output(
            tool_name="create_datatable",
            tool_args={},
            tool_output=output,
        )

        assert context is not None
        self.assertEqual(context["source_tool"], "create_datatable")
        self.assertEqual(context["data_connection"]["component_id"], "tool-1")
        self.assertEqual(context["columns"], ["Name", "Total"])
        self.assertEqual(context["rows"], [["Acme", 42]])

    def test_create_download_link_reuses_latest_export_context_when_source_omitted(self) -> None:
        export_context = {
            "latest": {
                "source_tool": "create_datatable",
                "title": "Accounts",
                "data_connection": {
                    "component_kind": "tool_config",
                    "component_id": "tool-1",
                    "request": {"query": "select name, total from accounts"},
                },
                "columns": ["Name", "Total"],
                "rows": [["Acme", 42]],
            }
        }
        tool = RAGComponents()._build_conversation_export_tool(
            conversation_id="conversation-1",
            workspace_id="workspace-1",
            export_context=export_context,
        )
        assert tool is not None and tool.coroutine is not None

        coroutine = cast(Coroutine[Any, Any, str], tool.coroutine(filename="accounts.xlsx", format="xlsx"))
        output: str = asyncio.run(coroutine)
        payload = json.loads(output)
        spec_path = next((export_service.EXPORT_BASE_DIR / "conversation-1").glob("*.json"))
        spec = json.loads(spec_path.read_text(encoding="utf-8"))

        self.assertTrue(payload["reused_previous_source"])
        self.assertEqual(payload["source_kind"], "live_table")
        self.assertEqual(spec["source"]["kind"], "live_table")
        self.assertEqual(spec["source"]["data_connection"]["component_id"], "tool-1")
        self.assertEqual(spec["source"]["snapshot_columns"], ["Name", "Total"])
        self.assertEqual(spec["source"]["snapshot_rows"], [["Acme", 42]])

    def test_seeds_latest_export_context_from_prior_chat_history(self) -> None:
        data_connection = {
            "component_kind": "tool_config",
            "component_id": "tool-1",
            "request": {"query": "select name, total from accounts"},
        }
        output = asyncio.run(
            create_datatable(
                title="Accounts",
                columns=["Name", "Total"],
                data=[["Acme", 42]],
                data_connection=data_connection,
            )
        )
        export_context: dict[str, Any] = {}

        RAGComponents()._seed_latest_export_context_from_chat_history(
            [
                AIMessage(
                    content="",
                    tool_calls=[
                        {
                            "name": "create_datatable",
                            "args": {"title": "Accounts"},
                            "id": "call_1",
                        }
                    ],
                ),
                ToolMessage(content=output, tool_call_id="call_1"),
            ],
            export_context,
        )

        self.assertEqual(export_context["latest"]["source_tool"], "create_datatable")
        self.assertEqual(export_context["latest"]["data_connection"]["component_id"], "tool-1")

    def test_extracts_latest_export_context_from_query_output(self) -> None:
        components = RAGComponents()
        components._tool_configs = [
            {
                "id": "tool-1",
                "name": "Warehouse DB",
                "tool_type": "postgres",
                "connection_config": {},
            }
        ]
        output = format_query_result(
            [{"name": "Acme", "total": 42}],
            ["name", "total"],
        )

        context = components._build_export_context_from_query_output(
            tool_name="query_warehouse_db",
            tool_args={"query": "select name, total from accounts limit 10"},
            tool_output=output,
        )

        assert context is not None
        self.assertEqual(context["source_tool"], "query_warehouse_db")
        self.assertEqual(context["data_connection"]["component_id"], "tool-1")
        self.assertEqual(context["data_connection"]["request"], {"query": "select name, total from accounts limit 10"})
        self.assertEqual(context["columns"], ["name", "total"])
        self.assertEqual(context["rows"], [["Acme", 42]])


class HtmlComponentToTableTests(unittest.TestCase):
    def test_dict_rows_are_mapped_by_column_order(self) -> None:
        payload = {
            "__html_component__": True,
            "title": "Shipments by origin",
            "html": "<!doctype html><html><body></body></html>",
            "data": {
                "columns": ["lat", "lng", "shipments"],
                # Row keys deliberately out of column order and one key missing.
                "rows": [
                    {"shipments": 412, "lat": 31.9, "lng": -99.9},
                    {"lat": 40.7, "shipments": 7},
                ],
                "row_count": 2,
            },
        }

        columns, rows = export_service.html_component_to_table(payload)

        self.assertEqual(columns, ["lat", "lng", "shipments"])
        self.assertEqual(rows, [[31.9, -99.9, 412], [40.7, "", 7]])

    def test_array_rows_pass_through_normalize_table(self) -> None:
        payload = {
            "__html_component__": True,
            "html": "<div/>",
            "data": {
                "columns": ["name", "total"],
                "rows": [["Acme", 42], ["Globex"], ["Initech", 1, "extra"]],
            },
        }

        columns, rows = export_service.html_component_to_table(payload)

        self.assertEqual(columns, ["name", "total"])
        self.assertEqual(rows, [["Acme", 42], ["Globex", ""], ["Initech", 1]])

    def test_non_tabular_data_raises_value_error(self) -> None:
        non_tabular: list[Any] = [
            None,
            {"lookup": {"TX": [31.9, -99.9]}},
            [{"lat": 1, "lng": 2}],
            {"columns": ["a"]},
            {"rows": [[1]]},
            {"columns": "a,b", "rows": [[1, 2]]},
            {"columns": ["a"], "rows": {"a": 1}},
        ]
        for data in non_tabular:
            with self.subTest(data=data):
                with self.assertRaises(ValueError):
                    export_service.html_component_to_table({"__html_component__": True, "html": "<div/>", "data": data})
        with self.assertRaises(ValueError):
            export_service.html_component_to_table({"__html_component__": True, "html": "<div/>"})


class HtmlComponentExportContextTests(unittest.TestCase):
    """Export context extraction for create_html_component tool output (PRD 6.7)."""

    _DATA_CONNECTION = {
        "component_kind": "tool_config",
        "component_id": "tool-1",
        "request": {"query": "select lat, lng, shipments from origins"},
    }

    @classmethod
    def _envelope(cls, data: Any, *, data_connection: dict[str, Any] | None = None) -> str:
        return json.dumps(
            {
                "__html_component__": True,
                "title": "Shipments by origin",
                "html": '<!doctype html><html><head></head><body><div id="map"></div></body></html>',
                "data": data,
                "description": "Shipments by origin state",
                "height": 480,
                "data_connection": cls._DATA_CONNECTION if data_connection is None else data_connection,
            }
        )

    def test_tabular_html_component_output_yields_export_context(self) -> None:
        output = self._envelope(
            {
                "columns": ["lat", "lng", "shipments"],
                "rows": [{"lat": 31.9, "lng": -99.9, "shipments": 412}],
                "row_count": 1,
            }
        )

        context = RAGComponents()._build_export_context_from_visualization_output(
            tool_name="create_html_component",
            tool_args={},
            tool_output=output,
        )

        assert context is not None
        self.assertEqual(context["source_tool"], "create_html_component")
        self.assertEqual(context["title"], "Shipments by origin")
        self.assertEqual(context["data_connection"]["component_id"], "tool-1")
        self.assertEqual(context["columns"], ["lat", "lng", "shipments"])
        self.assertEqual(context["rows"], [[31.9, -99.9, 412]])

    def test_non_tabular_html_component_output_yields_no_export_context(self) -> None:
        output = self._envelope({"nodes": [{"id": "a"}, {"id": "b"}], "edges": [["a", "b"]]})

        context = RAGComponents()._build_export_context_from_visualization_output(
            tool_name="create_html_component",
            tool_args={},
            tool_output=output,
        )

        self.assertIsNone(context)

    def test_html_component_output_without_data_connection_yields_no_export_context(self) -> None:
        output = self._envelope(
            {"columns": ["lat"], "rows": [{"lat": 31.9}], "row_count": 1},
            data_connection={},
        )

        context = RAGComponents()._build_export_context_from_visualization_output(
            tool_name="create_html_component",
            tool_args={},
            tool_output=output,
        )

        self.assertIsNone(context)


if __name__ == "__main__":
    unittest.main()
